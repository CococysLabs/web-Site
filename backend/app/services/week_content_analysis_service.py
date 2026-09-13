"""
Servicio de análisis con IA de:

- Semana Diagnóstico
- Semana 2 a Semana 11

El servicio:
1. Navega la estructura real de Google Drive.
2. Valida el código del curso usando el CourseCatalog recibido.
3. Calcula el período objetivo a partir del nombre real de la carpeta.
4. Comprueba los nombres y formatos de los archivos requeridos.
5. Lee F2_Semanas de 02_Matriz observaciones estructura.
6. Lee únicamente el contenido de las PRESENTACIONES.
7. Envía presentación + requisitos de la matriz a IA.
8. Escribe el resultado en Observaciones IA (columna H).

Los materiales Lectura, Ejemplo, Video y ActividadPractica son opcionales:
- no se analiza su contenido;
- si existen, se indican;
- si no existen, se deja vacío en la matriz.
"""

from __future__ import annotations

import io
import re
import time
import unicodedata
from typing import (
    Any,
    Callable,
    Dict,
    List,
    Optional,
    Tuple,
)

import openpyxl
from PyPDF2 import PdfReader

from app.config import settings
from app.models.course_catalog import CourseCatalog

from app.services.course_contacts_service import (
    course_contacts_service,
)

from app.services.curriculum_feedback_service import (
    curriculum_feedback_service,
)

from app.services.drive_service import (
    drive_service,
)

from app.services.google_sheets_service import (
    google_sheets_service,
)


# ============================================================
# MIME TYPES
# ============================================================

FOLDER_MIME = (
    "application/vnd.google-apps.folder"
)

GOOGLE_SHEET_MIME = (
    "application/vnd.google-apps.spreadsheet"
)

GOOGLE_DOC_MIME = (
    "application/vnd.google-apps.document"
)

GOOGLE_SLIDES_MIME = (
    "application/vnd.google-apps.presentation"
)

PDF_MIME = (
    "application/pdf"
)

DOCX_MIME = (
    "application/vnd.openxmlformats-officedocument."
    "wordprocessingml.document"
)

PPTX_MIME = (
    "application/vnd.openxmlformats-officedocument."
    "presentationml.presentation"
)

XLSX_MIME = (
    "application/vnd.openxmlformats-officedocument."
    "spreadsheetml.sheet"
)


# ============================================================
# ARCHIVOS QUE LA IA PUEDE LEER
# ============================================================

# IMPORTANTE:
# Tanto para Semana Diagnóstico como para Semana 2..11,
# el análisis de contenido se realiza EXCLUSIVAMENTE
# sobre archivos PDF.
ANALYSIS_PDF_MIMES = {
    PDF_MIME,
}


# Estos formatos pueden existir en Drive y pueden
# comprobarse como materiales, pero NO se leen con IA.
PRESENTATION_ASSET_MIMES = {
    PDF_MIME,
    PPTX_MIME,
    GOOGLE_SLIDES_MIME,
}

STRENGTHENING_MIMES = {
    PDF_MIME,
    DOCX_MIME,
    GOOGLE_DOC_MIME,
}


PRESENTATION_SHEET_MIMES = {
    GOOGLE_SHEET_MIME,
    XLSX_MIME,
}


MATRIX_MIMES = {
    GOOGLE_SHEET_MIME,
    XLSX_MIME,
}


# ============================================================
# CONFIGURACIÓN
# ============================================================

MATRIX_FOLDER_NAME = (
    "0_Revision_de_Material"
)

MATRIX_FILE_NAME = (
    "02_Matriz observaciones estructura"
)

MATRIX_SHEET_NAME = (
    "F2_Semanas"
)

MATRIX_AI_COLUMN = (
    "H"
)


DIAGNOSTIC_FOLDER_NAME = (
    "4_Diagnostico"
)

CONTENTS_FOLDER_NAME = (
    "5_Contenidos"
)


# Evitar enviar accidentalmente una presentación
# absurdamente grande a la IA.
MAX_PRESENTATION_CHARS = (
    220_000
)


class WeekContentAnalysisError(
    RuntimeError
):
    pass


class WeekContentAnalysisService:

    # ========================================================
    # NORMALIZACIÓN
    # ========================================================

    @staticmethod
    def normalize(
        value: object,
    ) -> str:

        text = str(
            value or ""
        ).strip().lower()

        text = re.sub(
            r"\.(pdf|pptx|docx|xlsx|xls|doc|ppt)$",
            "",
            text,
            flags=re.IGNORECASE,
        )

        text = text.replace(
            "_",
            " ",
        )

        text = text.replace(
            "-",
            " ",
        )

        text = text.replace(
            ".",
            " ",
        )

        normalized = (
            unicodedata.normalize(
                "NFKD",
                text,
            )
        )

        normalized = "".join(
            character
            for character in normalized
            if not unicodedata.combining(
                character
            )
        )

        normalized = re.sub(
            r"[^a-z0-9]+",
            " ",
            normalized,
        )

        return re.sub(
            r"\s+",
            " ",
            normalized,
        ).strip()


    @staticmethod
    def _file_stem(
        value: object,
    ) -> str:
        """
        Quitar únicamente una extensión conocida.

        Para nombres obligatorios queremos una
        comparación más estricta que normalize().
        """

        text = str(
            value or ""
        ).strip()

        return re.sub(
            r"\.(pdf|pptx|docx|xlsx|xls|doc|ppt)$",
            "",
            text,
            flags=re.IGNORECASE,
        )


    def _exact_file_name_match(
        self,
        actual_name: object,
        expected_base_name: str,
    ) -> bool:

        return (
            self._file_stem(
                actual_name
            ).casefold()
            ==
            str(
                expected_base_name
            ).strip().casefold()
        )


    @staticmethod
    def _cell(
        row: List[Any],
        index: int,
    ) -> str:

        if index >= len(
            row
        ):
            return ""

        value = row[
            index
        ]

        if value is None:
            return ""

        return str(
            value
        ).strip()


    # ========================================================
    # RAÍZ DE RECURSOS EDUCATIVOS
    # ========================================================

    def resources_root_folder_id(
        self,
    ) -> str:

        explicit = str(
            getattr(
                settings,
                "GOOGLE_DRIVE_RESOURCES_ROOT_FOLDER_ID",
                "",
            )
            or ""
        ).strip()

        if explicit:
            return explicit

        structure_root = str(
            getattr(
                settings,
                "GOOGLE_DRIVE_STRUCTURE_FOLDER_ID",
                "",
            )
            or ""
        ).strip()

        if not structure_root:

            raise WeekContentAnalysisError(
                "GOOGLE_DRIVE_RESOURCES_ROOT_FOLDER_ID "
                "no está configurado y tampoco existe "
                "GOOGLE_DRIVE_STRUCTURE_FOLDER_ID para "
                "obtener su carpeta padre"
            )

        parent = (
            drive_service
            .get_parent_folder_id(
                structure_root
            )
        )

        if not parent:

            raise WeekContentAnalysisError(
                "No se pudo resolver la carpeta raíz "
                "de Recursos Educativos"
            )

        return parent


    # ========================================================
    # NAVEGACIÓN DRIVE
    # ========================================================

    def list_semesters(
        self,
    ) -> Dict[str, Any]:
        """
        Devuelve las carpetas que existen realmente dentro de:

        GOOGLE_DRIVE_RESOURCES_ROOT_FOLDER_ID

        No hardcodea:
        - Primer Semestre
        - Segundo Semestre

        El frontend decidirá cuál selecciona el usuario.
        """

        root_id = (
            self.resources_root_folder_id()
        )

        root = (
            drive_service
            .get_file_metadata(
                root_id
            )
        )

        if not root:

            raise WeekContentAnalysisError(
                "No se pudo acceder a "
                "GOOGLE_DRIVE_RESOURCES_ROOT_FOLDER_ID"
            )

        folders = (
            drive_service
            .list_folders(
                root_id
            )
        )

        return {
            "success": True,

            "root": {
                "id": root_id,

                "name": (
                    root.get(
                        "name"
                    )
                ),

                "webViewLink": (
                    root.get(
                        "webViewLink"
                    )
                ),
            },

            "semesters": [
                {
                    "id": (
                        folder.get(
                            "id"
                        )
                    ),

                    "name": (
                        folder.get(
                            "name"
                        )
                    ),

                    "webViewLink": (
                        folder.get(
                            "webViewLink"
                        )
                    ),
                }

                for folder
                in folders
            ],
        }


    def list_areas(
        self,
        semester_folder_id: str,
    ) -> Dict[str, Any]:
        """
        Listar TODO lo que existe como carpeta dentro
        del semestre seleccionado.

        Ejemplo:

        2. Software
        3. Sistemas
        4. Computacion
        """

        semester = (
            drive_service
            .get_file_metadata(
                semester_folder_id
            )
        )

        if not semester:

            raise WeekContentAnalysisError(
                "No se pudo acceder a la carpeta "
                "de semestre seleccionada"
            )

        folders = (
            drive_service
            .list_folders(
                semester_folder_id
            )
        )

        return {
            "success": True,

            "semester": {
                "id": semester_folder_id,

                "name": (
                    semester.get(
                        "name"
                    )
                ),

                "webViewLink": (
                    semester.get(
                        "webViewLink"
                    )
                ),
            },

            "areas": [
                {
                    "id": (
                        folder.get(
                            "id"
                        )
                    ),

                    "name": (
                        folder.get(
                            "name"
                        )
                    ),

                    "canonical_area": (
                        course_contacts_service
                        .canonical_area(
                            folder.get(
                                "name",
                                "",
                            )
                        )
                    ),

                    "webViewLink": (
                        folder.get(
                            "webViewLink"
                        )
                    ),
                }

                for folder
                in folders
            ],
        }


    # ========================================================
    # NOMBRE REAL DE CARPETA DE CURSO
    # ========================================================

    def parse_course_folder_name(
        self,
        folder_name: str,
    ) -> Optional[
        Dict[str, Any]
    ]:
        """
        Ejemplo:

        92_Programacion_de_Computadoras_2_2S_2026

        resultado:

        code:
            92

        source:
            2S2026

        target:
            1S2027


        Otro ejemplo:

        XXX_Curso_1S_2027

        target:
            2S2027
        """

        name = str(
            folder_name
            or ""
        ).strip()

        match = re.match(
            r"^(?P<code>[^_]+)_"
            r"(?P<course_part>.+)_"
            r"(?P<semester>[12]S)_"
            r"(?P<year>\d{4})$",
            name,
            flags=re.IGNORECASE,
        )

        if not match:
            return None

        code = (
            match.group(
                "code"
            ).strip()
        )

        source_semester = (
            match.group(
                "semester"
            )
            .upper()
            .strip()
        )

        source_year = int(
            match.group(
                "year"
            )
        )

        # ----------------------------------------------------
        # CALCULAR SIGUIENTE SEMESTRE
        # ----------------------------------------------------

        if source_semester == "1S":

            target_semester = (
                "2S"
            )

            target_year = (
                source_year
            )

        else:

            target_semester = (
                "1S"
            )

            target_year = (
                source_year
                + 1
            )

        return {
            "folder_name": name,

            "code": code,

            "course_part": (
                match.group(
                    "course_part"
                )
            ),

            "source_semester": (
                source_semester
            ),

            "source_year": (
                source_year
            ),

            "target_semester": (
                target_semester
            ),

            "target_year": (
                target_year
            ),

            "source_period": (
                f"{source_semester}"
                f"{source_year}"
            ),

            "target_period": (
                f"{target_semester}"
                f"{target_year}"
            ),
        }


    def list_course_folders(
        self,
        area_folder_id: str,
    ) -> Dict[str, Any]:
        """
        Listar los cursos encontrados REALMENTE en Drive.

        Aquí todavía NO se consulta CourseCatalog.
        La ruta será la encargada de cruzar estos códigos
        contra la base de datos.
        """

        area = (
            drive_service
            .get_file_metadata(
                area_folder_id
            )
        )

        if not area:

            raise WeekContentAnalysisError(
                "No se pudo acceder a la carpeta "
                "de área seleccionada"
            )

        folders = (
            drive_service
            .list_folders(
                area_folder_id
            )
        )

        courses: List[
            Dict[str, Any]
        ] = []

        ignored: List[
            Dict[str, Any]
        ] = []

        for folder in folders:

            parsed = (
                self.parse_course_folder_name(
                    folder.get(
                        "name",
                        "",
                    )
                )
            )

            if not parsed:

                ignored.append(
                    {
                        "id": (
                            folder.get(
                                "id"
                            )
                        ),

                        "name": (
                            folder.get(
                                "name"
                            )
                        ),

                        "reason": (
                            "El nombre no coincide con "
                            "Codigo_Nombre_#S_Año"
                        ),
                    }
                )

                continue

            courses.append(
                {
                    "id": (
                        folder.get(
                            "id"
                        )
                    ),

                    "name": (
                        folder.get(
                            "name"
                        )
                    ),

                    "webViewLink": (
                        folder.get(
                            "webViewLink"
                        )
                    ),

                    **parsed,
                }
            )

        return {
            "success": True,

            "area": {
                "id": area_folder_id,

                "name": (
                    area.get(
                        "name"
                    )
                ),

                "canonical_area": (
                    course_contacts_service
                    .canonical_area(
                        area.get(
                            "name",
                            "",
                        )
                    )
                ),

                "webViewLink": (
                    area.get(
                        "webViewLink"
                    )
                ),
            },

            "courses": courses,

            "ignored": ignored,
        }


    # ========================================================
    # VALIDACIÓN DEL CURSO CONTRA BD
    # ========================================================

    def validate_course_identity(
        self,
        course: CourseCatalog,
        course_folder_name: str,
        area_folder_name: str,
    ) -> Dict[str, Any]:
        """
        El código de la carpeta REAL de Drive debe coincidir
        con CourseCatalog.

        No confiamos solamente en el nombre recibido por frontend.
        """

        parsed = (
            self.parse_course_folder_name(
                course_folder_name
            )
        )

        if not parsed:

            raise WeekContentAnalysisError(
                "El nombre de la carpeta del curso "
                "no tiene el formato esperado: "
                "Codigo_Nombre_#S_Año"
            )

        # ----------------------------------------------------
        # CÓDIGO
        # ----------------------------------------------------

        if (
            str(
                parsed[
                    "code"
                ]
            )
            !=
            str(
                course.code
            )
        ):

            raise WeekContentAnalysisError(
                "El código de la carpeta no coincide "
                "con el curso de la base de datos. "
                f"Carpeta={parsed['code']} | "
                f"BD={course.code}"
            )

        # ----------------------------------------------------
        # ÁREA
        # ----------------------------------------------------

        folder_area = (
            course_contacts_service
            .canonical_area(
                area_folder_name
            )
        )

        course_area = (
            course_contacts_service
            .canonical_area(
                course.area
            )
        )

        if (
            folder_area
            != course_area
        ):

            raise WeekContentAnalysisError(
                "El área de la carpeta no coincide "
                "con el área del curso en la base "
                "de datos. "
                f"Carpeta={folder_area} | "
                f"BD={course_area}"
            )

        return parsed


    # ========================================================
    # LOCALIZAR CARPETAS
    # ========================================================

    def _find_single_folder(
        self,
        parent_folder_id: str,
        expected_name: str,
        required: bool = False,
    ) -> Optional[
        Dict[str, Any]
    ]:

        folders = (
            drive_service
            .list_folders(
                parent_folder_id
            )
        )

        expected = (
            self.normalize(
                expected_name
            )
        )

        matches = [
            folder

            for folder
            in folders

            if (
                self.normalize(
                    folder.get(
                        "name",
                        "",
                    )
                )
                == expected
            )
        ]

        if len(
            matches
        ) > 1:

            raise WeekContentAnalysisError(
                "Se encontró más de una carpeta "
                f"equivalente a '{expected_name}'"
            )

        if matches:
            return matches[0]

        if required:

            raise WeekContentAnalysisError(
                "No se encontró la carpeta "
                f"'{expected_name}'"
            )

        return None


    # ========================================================
    # MATRIZ
    # ========================================================

    def _find_matrix_file(
        self,
        revision_folder_id: str,
    ) -> Optional[
        Dict[str, Any]
    ]:

        files = (
            drive_service
            .list_files(
                revision_folder_id
            )
        )

        matches = [
            file

            for file
            in files

            if (
                self.normalize(
                    file.get(
                        "name",
                        "",
                    )
                )
                ==
                self.normalize(
                    MATRIX_FILE_NAME
                )
            )
        ]

        if len(
            matches
        ) > 1:

            raise WeekContentAnalysisError(
                "Se encontró más de una matriz "
                f"llamada '{MATRIX_FILE_NAME}'"
            )

        if not matches:
            return None

        matrix = (
            matches[0]
        )

        if (
            matrix.get(
                "mimeType"
            )
            not in MATRIX_MIMES
        ):

            raise WeekContentAnalysisError(
                "02_Matriz observaciones estructura "
                "debe ser Google Sheets o XLSX"
            )

        return matrix


    # ========================================================
    # ARCHIVOS POR NOMBRE EXACTO
    # ========================================================

    def _matching_files(
        self,
        files: List[
            Dict[str, Any]
        ],
        expected_base_name: str,
    ) -> List[
        Dict[str, Any]
    ]:

        return [
            file

            for file
            in files

            if (
                self._exact_file_name_match(
                    file.get(
                        "name",
                        "",
                    ),
                    expected_base_name,
                )
            )
        ]


    @staticmethod
    def _filter_mimes(
        files: List[
            Dict[str, Any]
        ],
        allowed_mimes: set,
    ) -> List[
        Dict[str, Any]
    ]:

        return [
            file

            for file
            in files

            if (
                file.get(
                    "mimeType"
                )
                in allowed_mimes
            )
        ]


    @staticmethod
    def _file_public_data(
        file: Optional[
            Dict[str, Any]
        ],
    ) -> Optional[
        Dict[str, Any]
    ]:

        if not file:
            return None

        return {
            "id": (
                file.get(
                    "id"
                )
            ),

            "name": (
                file.get(
                    "name"
                )
            ),

            "mimeType": (
                file.get(
                    "mimeType"
                )
            ),

            "webViewLink": (
                file.get(
                    "webViewLink"
                )
            ),
        }


    def _select_analysis_pdf(
            self,
            candidates: List[
                Dict[str, Any]
            ],
        ) -> Optional[
            Dict[str, Any]
        ]:
        """
        Seleccionar EXCLUSIVAMENTE el PDF que será leído
        por la IA.

        Aunque exista:
        - Google Slides
        - PPTX
        - XLSX
        - cualquier otro archivo

        ninguno de ellos se utiliza para el análisis
        de contenido.

        El análisis de:
        - Semana Diagnóstico
        - Semana 2..11

        se realiza únicamente con PDF.
        """

        pdf_files = [
            file

            for file
            in candidates

            if (
                file.get(
                    "mimeType"
                )
                == PDF_MIME
            )
        ]

        if not pdf_files:
            return None

        if len(
            pdf_files
        ) > 1:

            raise WeekContentAnalysisError(
                "Se encontró más de un archivo PDF "
                "con el mismo nombre esperado. "
                "Debe existir un único PDF para "
                "realizar el análisis con IA."
            )

        return (
            pdf_files[0]
        )

    # ========================================================
    # LEER F2_SEMANAS — GOOGLE SHEETS
    # ========================================================

    def _matrix_google_sheet_values(
        self,
        spreadsheet_id: str,
    ) -> Tuple[
        str,
        List[List[Any]],
    ]:

        titles = (
            google_sheets_service
            .get_sheet_titles(
                spreadsheet_id
            )
        )

        matches = [
            title

            for title
            in titles

            if (
                self.normalize(
                    title
                )
                ==
                self.normalize(
                    MATRIX_SHEET_NAME
                )
            )
        ]

        if len(
            matches
        ) != 1:

            raise WeekContentAnalysisError(
                "La matriz debe contener exactamente "
                "una hoja 'F2_Semanas'"
            )

        title = (
            matches[0]
        )

        values = (
            google_sheets_service
            .get_sheet_values(
                spreadsheet_id,
                title,
            )
        )

        return (
            title,
            values,
        )


    # ========================================================
    # LEER F2_SEMANAS — XLSX
    # ========================================================

    def _matrix_xlsx_values(
        self,
        file_id: str,
    ) -> Tuple[
        str,
        List[List[Any]],
    ]:

        content = (
            drive_service
            .download_file(
                file_id
            )
        )

        if not content:

            raise WeekContentAnalysisError(
                "No se pudo descargar "
                "la matriz XLSX"
            )

        workbook = None

        try:

            workbook = (
                openpyxl
                .load_workbook(
                    io.BytesIO(
                        content
                    ),
                    read_only=True,
                    data_only=False,
                    keep_links=False,
                )
            )

            matches = [
                title

                for title
                in workbook.sheetnames

                if (
                    self.normalize(
                        title
                    )
                    ==
                    self.normalize(
                        MATRIX_SHEET_NAME
                    )
                )
            ]

            if len(
                matches
            ) != 1:

                raise WeekContentAnalysisError(
                    "La matriz XLSX debe contener "
                    "exactamente una hoja "
                    "'F2_Semanas'"
                )

            title = (
                matches[0]
            )

            worksheet = (
                workbook[
                    title
                ]
            )

            values: List[
                List[Any]
            ] = []

            for row in (
                worksheet
                .iter_rows(
                    min_row=1,
                    max_row=(
                        worksheet.max_row
                    ),
                    min_col=1,
                    max_col=8,
                    values_only=True,
                )
            ):

                values.append(
                    list(
                        row
                    )
                )

            return (
                title,
                values,
            )

        finally:

            if workbook is not None:
                workbook.close()


    def _read_matrix_values(
        self,
        matrix_file: Dict[
            str,
            Any,
        ],
    ) -> Tuple[
        str,
        List[List[Any]],
    ]:

        if (
            matrix_file.get(
                "mimeType"
            )
            == GOOGLE_SHEET_MIME
        ):

            return (
                self._matrix_google_sheet_values(
                    matrix_file[
                        "id"
                    ]
                )
            )

        if (
            matrix_file.get(
                "mimeType"
            )
            == XLSX_MIME
        ):

            return (
                self._matrix_xlsx_values(
                    matrix_file[
                        "id"
                    ]
                )
            )

        raise WeekContentAnalysisError(
            "Formato de matriz "
            "no soportado"
        )


    # ========================================================
    # IDENTIFICAR SEMANA EN F2_SEMANAS
    # ========================================================

    def _section_key(
        self,
        value: str,
    ) -> Optional[
        Tuple[
            str,
            Optional[int],
        ]
    ]:

        normalized = (
            self.normalize(
                value
            )
        )

        if normalized in {
            "semana diagnostico",
            "semana de diagnostico",
        }:

            return (
                "diagnostic",
                None,
            )

        match = re.fullmatch(
            r"semana\s*(\d+)",
            normalized,
        )

        if not match:
            return None

        week = int(
            match.group(
                1
            )
        )

        if (
            2
            <= week
            <= 11
        ):

            return (
                f"week_{week}",
                week,
            )

        return None


    # ========================================================
    # PARSEAR BLOQUES DE F2_SEMANAS
    # ========================================================

    def _parse_matrix_blocks(
        self,
        values: List[
            List[Any]
        ],
    ) -> Dict[
        str,
        Dict[str, Any],
    ]:
        """
        Lee dinámicamente:

        Columna A:
            Semana Diagnostico
            Semana 2
            ...
            Semana 11

        Columna B:
            Presentación
            Lectura
            Varios

        Columna C:
            requisitos

        NO utiliza números de fila hardcodeados.
        """

        blocks: Dict[
            str,
            Dict[str, Any],
        ] = {}

        current_key: Optional[
            str
        ] = None

        current_group: Optional[
            str
        ] = None


        for row_number, row in enumerate(
            values,
            start=1,
        ):

            section_value = (
                self._cell(
                    row,
                    0,
                )
            )

            group_value = (
                self._cell(
                    row,
                    1,
                )
            )

            requirement = (
                self._cell(
                    row,
                    2,
                )
            )

            # ------------------------------------------------
            # NUEVA SEMANA
            # ------------------------------------------------

            if section_value:

                section = (
                    self._section_key(
                        section_value
                    )
                )

                if section:

                    (
                        current_key,
                        week,
                    ) = section

                    current_group = (
                        None
                    )

                    if (
                        current_key
                        in blocks
                    ):

                        raise WeekContentAnalysisError(
                            "F2_Semanas contiene "
                            "una sección duplicada: "
                            f"{section_value}"
                        )

                    blocks[
                        current_key
                    ] = {
                        "key": (
                            current_key
                        ),

                        "label": (
                            section_value
                        ),

                        "week": (
                            week
                        ),

                        "section_row": (
                            row_number
                        ),

                        "presentation": {
                            "row": None,

                            "requirements": [],
                        },

                        "lectura": {
                            "row": None,
                        },

                        "varios": {
                            "video": None,

                            "actividad_practica": None,

                            "ejemplo": None,
                        },
                    }

                elif current_key:

                    current_key = (
                        None
                    )

                    current_group = (
                        None
                    )

            if not current_key:
                continue

            block = (
                blocks[
                    current_key
                ]
            )

            # ------------------------------------------------
            # GRUPO
            # ------------------------------------------------

            if group_value:

                current_group = (
                    self.normalize(
                        group_value
                    )
                )

                if (
                    current_group
                    == "presentacion"
                ):

                    if (
                        block[
                            "presentation"
                        ][
                            "row"
                        ]
                        is None
                    ):

                        block[
                            "presentation"
                        ][
                            "row"
                        ] = (
                            row_number
                        )

                elif (
                    current_group
                    == "lectura"
                ):

                    if (
                        block[
                            "lectura"
                        ][
                            "row"
                        ]
                        is None
                    ):

                        block[
                            "lectura"
                        ][
                            "row"
                        ] = (
                            row_number
                        )

            if not requirement:
                continue

            # ------------------------------------------------
            # PRESENTACIÓN
            # ------------------------------------------------

            if (
                current_group
                == "presentacion"
            ):

                block[
                    "presentation"
                ][
                    "requirements"
                ].append(
                    {
                        "row": (
                            row_number
                        ),

                        "label": (
                            requirement
                        ),
                    }
                )

            # ------------------------------------------------
            # VARIOS
            # ------------------------------------------------

            elif (
                current_group
                == "varios"
            ):

                normalized_requirement = (
                    self.normalize(
                        requirement
                    )
                )

                if (
                    "video"
                    in normalized_requirement
                ):

                    block[
                        "varios"
                    ][
                        "video"
                    ] = (
                        row_number
                    )

                elif (
                    "actividad"
                    in normalized_requirement
                ):

                    block[
                        "varios"
                    ][
                        "actividad_practica"
                    ] = (
                        row_number
                    )

                elif (
                    "ejemplo"
                    in normalized_requirement
                    or
                    "ejemlpo"
                    in normalized_requirement
                ):

                    block[
                        "varios"
                    ][
                        "ejemplo"
                    ] = (
                        row_number
                    )


        # ----------------------------------------------------
        # COMPROBAR QUE EXISTAN TODAS LAS SEMANAS
        # ----------------------------------------------------

        expected_keys = [
            "diagnostic",

            *[
                f"week_{week}"

                for week
                in range(
                    2,
                    12,
                )
            ],
        ]

        missing = [
            key

            for key
            in expected_keys

            if (
                key
                not in blocks
            )
        ]

        if missing:

            raise WeekContentAnalysisError(
                "F2_Semanas no contiene todas "
                "las secciones requeridas. "
                "Faltan: "
                + ", ".join(
                    missing
                )
            )

        # ----------------------------------------------------
        # TODAS DEBEN TENER PRESENTACIÓN
        # ----------------------------------------------------

        for key in expected_keys:

            presentation = (
                blocks[
                    key
                ][
                    "presentation"
                ]
            )

            if (
                presentation[
                    "row"
                ]
                is None
            ):

                raise WeekContentAnalysisError(
                    f"{blocks[key]['label']} "
                    "no contiene el bloque Presentación"
                )

            if not presentation[
                "requirements"
            ]:

                raise WeekContentAnalysisError(
                    f"{blocks[key]['label']} "
                    "no contiene requisitos "
                    "de Presentación"
                )

        return blocks


    def load_matrix(
        self,
        course_folder_id: str,
    ) -> Dict[str, Any]:

        revision_folder = (
            self._find_single_folder(
                course_folder_id,
                MATRIX_FOLDER_NAME,
                required=True,
            )
        )

        matrix_file = (
            self._find_matrix_file(
                revision_folder[
                    "id"
                ]
            )
        )

        if not matrix_file:

            raise WeekContentAnalysisError(
                "No se encontró "
                "'02_Matriz observaciones estructura'"
            )

        (
            sheet_title,
            values,
        ) = (
            self._read_matrix_values(
                matrix_file
            )
        )

        blocks = (
            self._parse_matrix_blocks(
                values
            )
        )

        return {
            "revision_folder": (
                revision_folder
            ),

            "file": (
                matrix_file
            ),

            "sheet_title": (
                sheet_title
            ),

            "blocks": (
                blocks
            ),
        }


    # ========================================================
    # NOMBRES ESPERADOS
    # ========================================================

    @staticmethod
    def expected_names(
        code: str,
        target_period: str,
        week: Optional[
            int
        ] = None,
    ) -> Dict[str, str]:

        # ----------------------------------------------------
        # DIAGNÓSTICO
        # ----------------------------------------------------

        if week is None:

            return {
                "diagnostico": (
                    f"{code}_Diagnostico_"
                    f"{target_period}"
                ),

                "t_fortalecimiento": (
                    f"{code}_TFortalecimiento_"
                    f"{target_period}"
                ),
            }

        # ----------------------------------------------------
        # SEMANAS
        # ----------------------------------------------------

        return {
            "presentacion": (
                f"{code}_Presentacion_"
                f"{week}_{target_period}"
            ),

            "lectura": (
                f"{code}_Lectura_"
                f"{week}_{target_period}"
            ),

            "ejemplo": (
                f"{code}_Ejemplo_"
                f"{week}_{target_period}"
            ),

            "video": (
                f"{code}_Video_"
                f"{week}_{target_period}"
            ),

            "actividad_practica": (
                f"{code}_ActividadPractica_"
                f"{week}_{target_period}"
            ),
        }


    # ========================================================
    # DIAGNÓSTICO
    # ========================================================

    def inspect_diagnostic(
        self,
        course_folder_id: str,
        code: str,
        target_period: str,
    ) -> Dict[str, Any]:

        folder = (
            self._find_single_folder(
                course_folder_id,
                DIAGNOSTIC_FOLDER_NAME,
                required=False,
            )
        )

        names = (
            self.expected_names(
                code,
                target_period,
            )
        )

        # ----------------------------------------------------
        # NO EXISTE 4_DIAGNOSTICO
        # ----------------------------------------------------

        if not folder:

            return {
                "folder": None,

                "expected": (
                    names
                ),

                "presentation": {
                    "found": False,

                    "valid_format": False,

                    "file": None,

                    "matching_files": [],
                },

                "strengthening": {
                    "found": False,

                    "valid_format": False,

                    "file": None,

                    "matching_files": [],
                },
            }

        files = (
            drive_service
            .list_files(
                folder[
                    "id"
                ]
            )
        )

        # ----------------------------------------------------
        # DIAGNÓSTICO
        #
        # Puede haber otros archivos con el mismo nombre base,
        # pero la IA únicamente utilizará el PDF.
        # ----------------------------------------------------

        diagnostic_matches = (
            self._matching_files(
                files,
                names[
                    "diagnostico"
                ],
            )
        )

        diagnostic_pdf = (
            self._select_analysis_pdf(
                diagnostic_matches
            )
        )

        # ----------------------------------------------------
        # TAREA DE FORTALECIMIENTO
        # ----------------------------------------------------

        strengthening_matches = (
            self._matching_files(
                files,
                names[
                    "t_fortalecimiento"
                ],
            )
        )

        strengthening_valid = (
            self._filter_mimes(
                strengthening_matches,
                STRENGTHENING_MIMES,
            )
        )

        strengthening_file = (
            strengthening_valid[
                0
            ]

            if strengthening_valid

            else None
        )

        return {
            "folder": (
                self._file_public_data(
                    folder
                )
            ),

            "expected": (
                names
            ),

            "presentation": {
                # Existe el material con ese nombre base.
                "found": bool(
                    diagnostic_matches
                ),

                # Para ANÁLISIS solamente importa que exista PDF.
                "analysis_pdf_found": bool(
                    diagnostic_pdf
                ),

                "valid_format": bool(
                    diagnostic_pdf
                ),

                # Este es EXCLUSIVAMENTE el PDF
                # que será enviado a extracción/IA.
                "file": (
                    self._file_public_data(
                        diagnostic_pdf
                    )
                ),

                # Se conservan únicamente como información.
                # NO se leen.
                "matching_files": [
                    self._file_public_data(
                        item
                    )

                    for item
                    in diagnostic_matches
                ],
            },

            "strengthening": {
                "found": bool(
                    strengthening_matches
                ),

                "valid_format": bool(
                    strengthening_file
                ),

                "file": (
                    self._file_public_data(
                        strengthening_file
                    )
                ),

                "matching_files": [
                    self._file_public_data(
                        item
                    )

                    for item
                    in strengthening_matches
                ],
            },
        }


    # ========================================================
    # SEMANA 2..11
    # ========================================================

    def inspect_week(
        self,
        contents_folder: Optional[
            Dict[str, Any]
        ],
        code: str,
        target_period: str,
        week: int,
    ) -> Dict[str, Any]:

        names = (
            self.expected_names(
                code,
                target_period,
                week=week,
            )
        )

        empty_result = {
            "week": (
                week
            ),

            "folder": None,

            "expected": (
                names
            ),

            "presentation": {
                "found": False,

                "valid_format": False,

                "file": None,

                "matching_files": [],
            },

            "presentation_sheet": {
                "found": False,

                "valid_format": False,

                "file": None,

                "matching_files": [],
            },

            "optional": {
                "lectura": [],

                "ejemplo": [],

                "video": [],

                "actividad_practica": [],
            },
        }

        if not contents_folder:
            return empty_result

        # ----------------------------------------------------
        # EJEMPLO:
        # 2_Semana_2
        # ----------------------------------------------------

        folder = (
            self._find_single_folder(
                contents_folder[
                    "id"
                ],
                f"{week}_Semana_{week}",
                required=False,
            )
        )

        if not folder:
            return empty_result

        files = (
            drive_service
            .list_files(
                folder[
                    "id"
                ]
            )
        )

        # ----------------------------------------------------
        # PRESENTACIÓN DE LA SEMANA
        #
        # Buscamos todos los archivos con el nombre esperado
        # porque también puede existir el Google Sheets/XLSX.
        #
        # Pero para IA seleccionamos EXCLUSIVAMENTE PDF.
        # ----------------------------------------------------

        presentation_matches = (
            self._matching_files(
                files,
                names[
                    "presentacion"
                ],
            )
        )

        presentation_pdf = (
            self._select_analysis_pdf(
                presentation_matches
            )
        )

        # ----------------------------------------------------
        # GOOGLE SHEETS/XLSX
        #
        # MISMO NOMBRE DE LA PRESENTACIÓN
        # ----------------------------------------------------

        sheet_matches = (
            self._filter_mimes(
                presentation_matches,
                PRESENTATION_SHEET_MIMES,
            )
        )

        presentation_sheet = (
            sheet_matches[
                0
            ]

            if sheet_matches

            else None
        )

        # ----------------------------------------------------
        # OPCIONALES
        # ----------------------------------------------------

        optional: Dict[
            str,
            List[
                Dict[str, Any]
            ],
        ] = {}

        for key in (
            "lectura",
            "ejemplo",
            "video",
            "actividad_practica",
        ):

            matches = (
                self._matching_files(
                    files,
                    names[
                        key
                    ],
                )
            )

            optional[
                key
            ] = [
                self._file_public_data(
                    item
                )

                for item
                in matches
            ]

        return {
            "week": (
                week
            ),

            "folder": (
                self._file_public_data(
                    folder
                )
            ),

            "expected": (
                names
            ),

            "presentation": {
                "found": bool(
                    presentation_matches
                ),

                # Lo que realmente interesa para IA.
                "analysis_pdf_found": bool(
                    presentation_pdf
                ),

                "valid_format": bool(
                    presentation_pdf
                ),

                # EXCLUSIVAMENTE PDF.
                "file": (
                    self._file_public_data(
                        presentation_pdf
                    )
                ),

                # Solo informativo.
                "matching_files": [
                    self._file_public_data(
                        item
                    )

                    for item
                    in presentation_matches
                ],
            },

            "presentation_sheet": {
                "found": bool(
                    sheet_matches
                ),

                "valid_format": bool(
                    presentation_sheet
                ),

                "file": (
                    self._file_public_data(
                        presentation_sheet
                    )
                ),

                "matching_files": [
                    self._file_public_data(
                        item
                    )

                    for item
                    in sheet_matches
                ],
            },

            "optional": (
                optional
            ),
        }


    # ========================================================
    # PREVIEW
    # ========================================================

    def preview_course(
        self,
        course: CourseCatalog,
        course_folder_id: str,
        course_folder_name: str,
        area_folder_name: str,
    ) -> Dict[str, Any]:
        """
        Preview:

        - NO llama IA.
        - NO modifica la matriz.
        - Sí comprueba estructura.
        - Sí comprueba nombres.
        - Sí comprueba formatos.
        """

        try:

            period = (
                self.validate_course_identity(
                    course,
                    course_folder_name,
                    area_folder_name,
                )
            )

            matrix = (
                self.load_matrix(
                    course_folder_id
                )
            )

            diagnostic = (
                self.inspect_diagnostic(
                    course_folder_id,
                    str(
                        course.code
                    ),
                    period[
                        "target_period"
                    ],
                )
            )

            contents_folder = (
                self._find_single_folder(
                    course_folder_id,
                    CONTENTS_FOLDER_NAME,
                    required=False,
                )
            )

            weeks = {
                str(
                    week
                ): (
                    self.inspect_week(
                        contents_folder,
                        str(
                            course.code
                        ),
                        period[
                            "target_period"
                        ],
                        week,
                    )
                )

                for week
                in range(
                    2,
                    12,
                )
            }

            warnings: List[
                str
            ] = []

            # ------------------------------------------------
            # DIAGNÓSTICO
            # ------------------------------------------------

            if not diagnostic[
                "folder"
            ]:

                warnings.append(
                    "No se encontró "
                    "4_Diagnostico"
                )

            if not diagnostic[
                "presentation"
            ][
                "valid_format"
            ]:

                warnings.append(
                    "No se encontró el PDF de Diagnóstico "
                    "requerido para realizar el análisis con IA"
                )

            if not diagnostic[
                "strengthening"
            ][
                "valid_format"
            ]:

                warnings.append(
                    "No se encontró "
                    "TFortalecimiento obligatorio "
                    "con formato "
                    "PDF/DOCX/Google Docs"
                )

            # ------------------------------------------------
            # CONTENIDOS
            # ------------------------------------------------

            if not contents_folder:

                warnings.append(
                    "No se encontró "
                    "5_Contenidos"
                )

            # ------------------------------------------------
            # SEMANAS
            # ------------------------------------------------

            for week in range(
                2,
                12,
            ):

                item = (
                    weeks[
                        str(
                            week
                        )
                    ]
                )

                if not item[
                    "folder"
                ]:

                    warnings.append(
                        "No se encontró "
                        f"{week}_Semana_{week}"
                    )

                    continue

                if not item[
                    "presentation"
                ][
                    "valid_format"
                ]:

                    warnings.append(
                        f"Semana {week}: "
                        "no se encontró el PDF de la presentación "
                        "requerido para realizar el análisis con IA"
                    )

                if not item[
                    "presentation_sheet"
                ][
                    "valid_format"
                ]:

                    warnings.append(
                        f"Semana {week}: "
                        "no se encontró el "
                        "Google Sheets/XLSX obligatorio "
                        "con el mismo nombre "
                        "de la presentación"
                    )

            return {
                "success": True,

                # Aunque falten materiales, el curso
                # puede analizarse para que la IA
                # REPORTE justamente esas ausencias.
                "ready_for_analysis": True,

                "ready_for_write": True,

                "course": {
                    "code": str(
                        course.code
                    ),

                    "name": (
                        course.name
                    ),

                    "area": (
                        course_contacts_service
                        .canonical_area(
                            course.area
                        )
                    ),

                    "folder_id": (
                        course_folder_id
                    ),

                    "folder_name": (
                        course_folder_name
                    ),
                },

                "period": (
                    period
                ),

                "matrix": {
                    "file": (
                        self._file_public_data(
                            matrix[
                                "file"
                            ]
                        )
                    ),

                    "sheet": (
                        matrix[
                            "sheet_title"
                        ]
                    ),

                    "column": (
                        MATRIX_AI_COLUMN
                    ),
                },

                "diagnostic": (
                    diagnostic
                ),

                "contents_folder": (
                    self._file_public_data(
                        contents_folder
                    )
                ),

                "weeks": (
                    weeks
                ),

                "warnings": (
                    warnings
                ),

                "error": None,
            }

        except Exception as exc:

            return {
                "success": False,

                "ready_for_analysis": False,

                "ready_for_write": False,

                "course": {
                    "code": str(
                        course.code
                    ),

                    "name": (
                        course.name
                    ),

                    "area": (
                        course.area
                    ),

                    "folder_id": (
                        course_folder_id
                    ),

                    "folder_name": (
                        course_folder_name
                    ),
                },

                "warnings": [],

                "error": str(
                    exc
                ),
            }


    # ========================================================
    # PDF -> TEXTO
    # ========================================================

    @staticmethod
    def _pdf_to_text(
        content: bytes,
    ) -> str:

        reader = (
            PdfReader(
                io.BytesIO(
                    content
                )
            )
        )

        parts: List[
            str
        ] = []

        for index, page in enumerate(
            reader.pages,
            start=1,
        ):

            text = (
                page.extract_text()
                or ""
            ).strip()

            if text:

                parts.append(
                    f"=== Página {index} ===\n"
                    f"{text}"
                )

        return "\n\n".join(
            parts
        ).strip()

    # ========================================================
    # EXTRAER PRESENTACIÓN
    # ========================================================

    def extract_presentation_text(
            self,
            file: Optional[
                Dict[str, Any]
            ],
        ) -> Dict[str, Any]:
        """
        Extraer contenido para análisis con IA.

        REGLA OBLIGATORIA:

        Tanto Semana Diagnóstico como Semana 2..11
        se leen EXCLUSIVAMENTE desde PDF.

        NO se procesa:

        - PPTX
        - Google Slides
        - DOCX
        - Google Docs
        - XLSX
        - Google Sheets
        - Lectura
        - Video
        - Ejemplo
        - Actividad práctica
        """

        if not file:

            return {
                "success": False,

                "text": "",

                "chars": 0,

                "original_chars": 0,

                "truncated": False,

                "error": (
                    "No se encontró el PDF "
                    "requerido para realizar "
                    "el análisis con IA"
                ),
            }

        mime_type = str(
            file.get(
                "mimeType"
            )
            or ""
        ).strip()

        # ----------------------------------------------------
        # SOLO PDF
        # ----------------------------------------------------

        if (
            mime_type
            != PDF_MIME
        ):

            return {
                "success": False,

                "text": "",

                "chars": 0,

                "original_chars": 0,

                "truncated": False,

                "error": (
                    "El archivo encontrado no es PDF. "
                    "Este módulo únicamente analiza "
                    "contenido proveniente de archivos PDF."
                ),
            }

        # ----------------------------------------------------
        # DESCARGAR PDF
        # ----------------------------------------------------

        content = (
            drive_service
            .download_file(
                file[
                    "id"
                ]
            )
        )

        if not content:

            return {
                "success": False,

                "text": "",

                "chars": 0,

                "original_chars": 0,

                "truncated": False,

                "error": (
                    "No se pudo descargar "
                    "el archivo PDF"
                ),
            }

        # ----------------------------------------------------
        # EXTRAER TEXTO
        # ----------------------------------------------------

        try:

            text = (
                self._pdf_to_text(
                    content
                )
            )

        except Exception as exc:

            return {
                "success": False,

                "text": "",

                "chars": 0,

                "original_chars": 0,

                "truncated": False,

                "error": (
                    "No se pudo leer "
                    f"el archivo PDF: {exc}"
                ),
            }

        original_chars = (
            len(
                text
            )
        )

        truncated = False

        # ----------------------------------------------------
        # LÍMITE DE SEGURIDAD
        # ----------------------------------------------------

        if (
            len(
                text
            )
            > MAX_PRESENTATION_CHARS
        ):

            text = (
                text[
                    :MAX_PRESENTATION_CHARS
                ]
            )

            truncated = True

        # ----------------------------------------------------
        # PDF SIN TEXTO EXTRAÍBLE
        # ----------------------------------------------------

        if not text.strip():

            return {
                "success": False,

                "text": "",

                "chars": 0,

                "original_chars": (
                    original_chars
                ),

                "truncated": (
                    truncated
                ),

                "error": (
                    "El PDF existe, pero no produjo "
                    "texto utilizable para el análisis. "
                    "Puede estar compuesto únicamente "
                    "por imágenes."
                ),
            }

        # ----------------------------------------------------
        # OK
        # ----------------------------------------------------

        return {
            "success": True,

            "text": (
                text
            ),

            "chars": (
                len(
                    text
                )
            ),

            "original_chars": (
                original_chars
            ),

            "truncated": (
                truncated
            ),

            "error": None,
        }

    @staticmethod
    def _build_observation_from_criteria(
        criteria: List[
            Dict[str, Any]
        ],
        sheet_found: Optional[
            bool
        ] = None,
    ) -> str:
        """
        Construir una observación legible a partir de las
        decisiones estructuradas de la IA.

        Sirve además como respaldo si el proveedor devuelve
        correctamente los criterios pero omite
        observacion_presentacion.
        """

        present = [
            item[
                "criterio"
            ]

            for item
            in criteria

            if (
                item.get(
                    "estado"
                )
                == "PRESENTE"
            )
        ]

        missing = [
            item[
                "criterio"
            ]

            for item
            in criteria

            if (
                item.get(
                    "estado"
                )
                == "NO_ENCONTRADO"
            )
        ]

        not_evaluable = [
            item[
                "criterio"
            ]

            for item
            in criteria

            if (
                item.get(
                    "estado"
                )
                == "NO_EVALUABLE"
            )
        ]

        parts: List[
            str
        ] = []

        if present:

            parts.append(
                "Presentes:\n"
                + "\n".join(
                    f"- {item}"
                    for item
                    in present
                )
            )

        if missing:

            parts.append(
                "No encontrados:\n"
                + "\n".join(
                    f"- {item}"
                    for item
                    in missing
                )
            )

        if not_evaluable:

            parts.append(
                "No evaluables:\n"
                + "\n".join(
                    f"- {item}"
                    for item
                    in not_evaluable
                )
            )

        if (
            sheet_found
            is not None
        ):

            parts.append(
                "Archivo Google Sheets/XLSX:\n"
                + (
                    "- Encontrado"
                    if sheet_found
                    else "- No encontrado"
                )
            )

        return (
            "\n\n".join(
                parts
            )
        ).strip()


    # ========================================================
    # ESQUEMA DE RESPUESTA IA
    # ========================================================

    def _ai_response_schema(
        self,
        requirements: List[
            Dict[str, Any]
        ],
    ) -> Dict[str, Any]:
        """
        El criterio se identifica principalmente por su ÍNDICE,
        no porque la IA repita literalmente el texto.

        El backend sabe que el criterio 1 corresponde al primer
        requisito recibido.
        """

        return {
            "type": "OBJECT",

            "properties": {
                "criterios": {
                    "type": "ARRAY",

                    "items": {
                        "type": "OBJECT",

                        "properties": {
                            "indice": {
                                "type": "INTEGER",
                            },

                            "criterio": {
                                "type": "STRING",
                            },

                            "estado": {
                                "type": "STRING",

                                "enum": [
                                    "PRESENTE",
                                    "NO_ENCONTRADO",
                                    "NO_EVALUABLE",
                                ],
                            },

                            "evidencia": {
                                "type": "STRING",
                            },
                        },

                        "required": [
                            "indice",
                            "criterio",
                            "estado",
                            "evidencia",
                        ],
                    },
                },

                "observacion_presentacion": {
                    "type": "STRING",
                },

                "lectura": {
                    "type": "STRING",
                },

                "video": {
                    "type": "STRING",
                },

                "actividad_practica": {
                    "type": "STRING",
                },

                "ejemplo": {
                    "type": "STRING",
                },
            },

            "required": [
                "criterios",
                "observacion_presentacion",
                "lectura",
                "video",
                "actividad_practica",
                "ejemplo",
            ],
        }


    # ========================================================
    # PROMPT IA
    # ========================================================

    def _build_ai_prompt(
        self,
        course: CourseCatalog,
        section_label: str,
        requirements: List[
            Dict[str, Any]
        ],
        presentation_info: Dict[
            str,
            Any,
        ],
        presentation_sheet_info: Optional[
            Dict[str, Any]
        ],
        extraction: Dict[
            str,
            Any,
        ],
        optional_files: Optional[
            Dict[
                str,
                List[
                    Dict[str, Any]
                ],
            ]
        ] = None,
        strengthening_info: Optional[
            Dict[str, Any]
        ] = None,
    ) -> Tuple[
        str,
        str,
    ]:

        # ====================================================
        # SYSTEM
        # ====================================================

        system_message = """
Eres un revisor de estructura de materiales educativos universitarios.

Tu tarea en este módulo NO es evaluar la calidad pedagógica,
la coherencia curricular ni hacer un resumen.

Tu tarea es comprobar, utilizando IA, si cada elemento exigido
por la matriz F2_Semanas está realmente presente en el contenido
extraído del ARCHIVO PDF de la presentación.

IMPORTANTE:

El backend únicamente proporciona a la IA contenido proveniente
de archivos PDF.

No se analiza el contenido de:

- PPTX;
- Google Slides;
- Google Sheets;
- XLSX;
- Lectura;
- Video;
- Actividad práctica;
- Ejemplo demostrativo.

============================================================
REGLAS
============================================================

1. Evalúa semánticamente.

Un requisito puede estar presente aunque el título literal no
aparezca, siempre que exista contenido REAL que cumpla esa función.

Ejemplo:

Si el requisito es "Bienvenida", una diapositiva que diga
"Bienvenidos al curso..." puede cumplir aunque no tenga el título
literal "Bienvenida".


2. NO marques PRESENTE un requisito solamente porque su nombre
aparece mencionado en:

- una agenda;
- un índice;
- una lista;
- una tabla de contenidos;
- una frase que únicamente anuncia que después se verá.

Debe existir contenido real correspondiente.


3. NO inventes contenido.


4. HECHOS DE ARCHIVOS

La existencia de archivos NO debe inferirse de la presentación.

Debes usar EXCLUSIVAMENTE los HECHOS DEL BACKEND.


5. REQUISITOS DE ARCHIVO

Cuando el requisito sea equivalente a:

"Presentación pptx/canva y pdf"

o:

"Archivo pptx/canva y pdf"

NO debes buscarlo dentro del texto.

Su estado depende EXCLUSIVAMENTE del hecho:

PRESENTACIÓN OBLIGATORIA.


Cuando el requisito sea:

"Tarea de Fortalecimiento"

su existencia depende EXCLUSIVAMENTE del hecho:

TAREA DE FORTALECIMIENTO.


6. PDF SIN TEXTO EXTRAÍBLE

Si el PDF existe pero el backend no pudo obtener texto utilizable:

- los requisitos relacionados exclusivamente con existencia
  de archivos sí pueden determinarse mediante los hechos del backend;
- los requisitos relacionados con contenido deben marcarse
  NO_EVALUABLE.

No concluyas que el contenido está ausente cuando el PDF
simplemente no pudo ser leído.


7. GOOGLE SHEETS/XLSX

Para Semana 2 a Semana 11 debe existir además un Google Sheets
o XLSX con EXACTAMENTE el mismo nombre base de la presentación.

Ese archivo:

- es obligatorio;
- NO debe analizarse;
- solo debes indicar en observacion_presentacion si fue encontrado
  o no.


8. MATERIALES OPCIONALES

Lectura
Video
Actividad práctica
Ejemplo demostrativo

son opcionales.

NO analices su contenido.

Si el backend dice que existe:

indica únicamente que fue encontrado y menciona su nombre.

Si no existe:

devuelve exactamente "" para ese campo.

NO lo reportes como error.


9. OBSERVACIÓN DE PRESENTACIÓN

El campo observacion_presentacion se escribirá DIRECTAMENTE
en la columna Observaciones IA.

Debe indicar claramente:

Presentes:
- requisito
- requisito

No encontrados:
- requisito
- requisito

No evaluables:
- requisito

Solo incluye "No evaluables" cuando aplique.

Para las semanas normales añade también:

Archivo Google Sheets/XLSX:
- Encontrado

o:

Archivo Google Sheets/XLSX:
- No encontrado


10. NO menciones:

- filas;
- celdas;
- coordenadas;
- nombres internos del backend;
- JSON;
- análisis curricular;
- aspectos que no pertenezcan a esta presentación.


11. Devuelve TODOS los requisitos recibidos exactamente una vez.

No agregues requisitos nuevos.

No elimines requisitos.

Responde únicamente JSON válido.
""".strip()


        # ====================================================
        # REQUISITOS
        # ====================================================

        requirement_text = "\n".join(
            (
                f"{index}. "
                f"{item['label']}"
            )

            for index, item
            in enumerate(
                requirements,
                start=1,
            )
        )


        # ====================================================
        # PRESENTACIÓN
        # ====================================================

        presentation_file = (
            presentation_info.get(
                "file"
            )
        )

        if presentation_info.get(
            "valid_format"
        ):

            presentation_status = (
                "ENCONTRADA Y CON FORMATO VÁLIDO"
            )

        elif presentation_info.get(
            "found"
        ):

            presentation_status = (
                "ENCONTRADA PERO CON "
                "FORMATO NO VÁLIDO"
            )

        else:

            presentation_status = (
                "NO ENCONTRADA"
            )


        presentation_name = (
            presentation_file.get(
                "name",
                "",
            )

            if presentation_file

            else ""
        )


        # ====================================================
        # SHEET/XLSX
        # ====================================================

        if (
            presentation_sheet_info
            is None
        ):

            sheet_status = (
                "NO APLICA"
            )

            sheet_name = ""

        else:

            if presentation_sheet_info.get(
                "valid_format"
            ):

                sheet_status = (
                    "ENCONTRADO"
                )

            else:

                sheet_status = (
                    "NO ENCONTRADO"
                )

            sheet_file = (
                presentation_sheet_info.get(
                    "file"
                )
            )

            sheet_name = (
                sheet_file.get(
                    "name",
                    "",
                )

                if sheet_file

                else ""
            )


        # ====================================================
        # TAREA FORTALECIMIENTO
        # ====================================================

        strengthening_text = (
            "NO APLICA"
        )

        if (
            strengthening_info
            is not None
        ):

            strengthening_file = (
                strengthening_info.get(
                    "file"
                )
            )

            if strengthening_info.get(
                "valid_format"
            ):

                strengthening_text = (
                    "ENCONTRADA Y CON FORMATO VÁLIDO"
                )

            elif strengthening_info.get(
                "found"
            ):

                strengthening_text = (
                    "ENCONTRADA PERO CON "
                    "FORMATO NO VÁLIDO"
                )

            else:

                strengthening_text = (
                    "NO ENCONTRADA"
                )

            if strengthening_file:

                strengthening_text += (
                    " | archivo="
                    + str(
                        strengthening_file.get(
                            "name",
                            "",
                        )
                    )
                )


        # ====================================================
        # OPCIONALES
        # ====================================================

        optional_files = (
            optional_files
            or {}
        )

        optional_lines: List[
            str
        ] = []

        for key, label in (
            (
                "lectura",
                "Lectura",
            ),
            (
                "video",
                "Video",
            ),
            (
                "actividad_practica",
                "Actividad práctica",
            ),
            (
                "ejemplo",
                "Ejemplo demostrativo",
            ),
        ):

            found = (
                optional_files.get(
                    key,
                    [],
                )
            )

            names = [
                str(
                    item.get(
                        "name"
                    )
                    or ""
                )

                for item
                in found

                if item
            ]

            if names:

                optional_lines.append(
                    f"- {label}: "
                    "ENCONTRADO | "
                    + ", ".join(
                        names
                    )
                )

            else:

                optional_lines.append(
                    f"- {label}: "
                    "NO ENCONTRADO"
                )


        # ====================================================
        # EXTRACCIÓN
        # ====================================================

        if extraction.get(
            "success"
        ):

            extraction_status = (
                "TEXTO EXTRAÍDO CORRECTAMENTE"
            )

        else:

            extraction_status = (
                "NO SE PUDO OBTENER TEXTO "
                "UTILIZABLE: "
                + str(
                    extraction.get(
                        "error"
                    )
                    or "motivo desconocido"
                )
            )

        presentation_text = str(
            extraction.get(
                "text"
            )
            or ""
        )


        # ====================================================
        # USER
        # ====================================================

        user_message = f"""
============================================================
CURSO
============================================================

Código:
{course.code}

Nombre:
{course.name}


============================================================
SECCIÓN DE LA MATRIZ
============================================================

{section_label}


============================================================
REQUISITOS QUE DEBES EVALUAR
============================================================

{requirement_text}


============================================================
HECHOS DEL BACKEND — SON AUTORITATIVOS
============================================================

PRESENTACIÓN OBLIGATORIA

Estado:
{presentation_status}

Archivo:
{presentation_name or "Ninguno"}


GOOGLE SHEETS/XLSX OBLIGATORIO
CON EL MISMO NOMBRE DE LA PRESENTACIÓN

Estado:
{sheet_status}

Archivo:
{sheet_name or "Ninguno"}


TAREA DE FORTALECIMIENTO

Estado:
{strengthening_text}


MATERIALES OPCIONALES

{chr(10).join(optional_lines)}


============================================================
EXTRACCIÓN DE LA PRESENTACIÓN
============================================================

Estado:
{extraction_status}

Texto truncado por seguridad del backend:
{"Sí" if extraction.get("truncated") else "No"}


============================================================
CONTENIDO EXTRAÍDO DE LA PRESENTACIÓN
============================================================

{presentation_text}


============================================================
SALIDA
============================================================

Devuelve TODOS los requisitos exactamente una vez dentro
de "criterios".

MUY IMPORTANTE:

Cada requisito tiene un número al inicio.

Debes devolver ese número exactamente dentro de "indice".

Ejemplo:

Si recibes:

1. Bienvenida
2. Agenda
3. Competencias

debes devolver:

{{
    "indice": 1,
    "criterio": "Bienvenida",
    "estado": "PRESENTE",
    "evidencia": "..."
}}

{{
    "indice": 2,
    "criterio": "Agenda",
    "estado": "PRESENTE",
    "evidencia": "..."
}}

y así sucesivamente.

El campo "indice" es OBLIGATORIO.

No cambies el orden lógico de los requisitos.

Para cada requisito:

indice:
número exacto recibido.

criterio:
texto del requisito.

estado:
PRESENTE
NO_ENCONTRADO
NO_EVALUABLE

evidencia:
explicación breve del contenido que permitió tomar
la decisión.

Después genera:

observacion_presentacion

que será el texto que se escribirá directamente
en Observaciones IA.

Para:

lectura
video
actividad_practica
ejemplo

si el archivo NO existe:

""

Si existe:

indica únicamente que fue encontrado y el nombre del archivo.

No analices su contenido.
""".strip()

        return (
            system_message,
            user_message,
        )


    # ========================================================
    # VALIDAR RESPUESTA IA
    # ========================================================

    def _validate_ai_payload(
        self,
        payload: Dict[str, Any],
        requirements: List[
            Dict[str, Any]
        ],
        optional_files: Optional[
            Dict[
                str,
                List[
                    Dict[str, Any]
                ],
            ]
        ] = None,
    ) -> Dict[str, Any]:
        """
        Validar la respuesta de IA.

        IMPORTANTE:

        La asociación principal se realiza por ÍNDICE,
        NO por coincidencia literal del nombre del criterio.

        Esto evita que una pequeña variación de redacción de la IA
        provoque que toda una semana sea considerada inválida.
        """

        criteria = (
            payload.get(
                "criterios"
            )
        )

        if not isinstance(
            criteria,
            list,
        ):

            raise WeekContentAnalysisError(
                "La IA no devolvió "
                "el arreglo 'criterios'"
            )

        if not criteria:

            raise WeekContentAnalysisError(
                "La IA devolvió el arreglo "
                "'criterios' vacío"
            )

        allowed_states = {
            "PRESENTE",
            "NO_ENCONTRADO",
            "NO_EVALUABLE",
        }

        total_requirements = (
            len(
                requirements
            )
        )

        received_by_index: Dict[
            int,
            Dict[str, str],
        ] = {}

        # ====================================================
        # 1. PROCESAR RESPUESTA
        # ====================================================

        for position, item in enumerate(
            criteria,
            start=1,
        ):

            if not isinstance(
                item,
                dict,
            ):
                continue

            # ------------------------------------------------
            # ÍNDICE
            #
            # Compatibilidad:
            #
            # Si la IA olvidó "indice", pero devolvió exactamente
            # la misma cantidad y en el mismo orden, podemos usar
            # la posición del elemento.
            # ------------------------------------------------

            raw_index = (
                item.get(
                    "indice"
                )
            )

            try:

                criterion_index = int(
                    raw_index
                )

            except (
                TypeError,
                ValueError,
            ):

                criterion_index = (
                    position
                )

            if (
                criterion_index < 1
                or
                criterion_index
                > total_requirements
            ):

                # Ignorar índices inventados.
                continue

            if (
                criterion_index
                in received_by_index
            ):

                raise WeekContentAnalysisError(
                    "La IA devolvió dos veces "
                    "el criterio con índice "
                    f"{criterion_index}"
                )

            # ------------------------------------------------
            # ESTADO
            # ------------------------------------------------

            state = str(
                item.get(
                    "estado"
                )
                or ""
            ).strip().upper()

            if (
                state
                not in allowed_states
            ):

                raise WeekContentAnalysisError(
                    "La IA devolvió un estado "
                    "inválido para el criterio "
                    f"{criterion_index}: "
                    f"{state}"
                )

            # ------------------------------------------------
            # NOMBRE ORIGINAL
            #
            # NO confiamos en el nombre retornado por IA.
            #
            # Utilizamos SIEMPRE el nombre original de F2_Semanas.
            # ------------------------------------------------

            original_requirement = (
                requirements[
                    criterion_index - 1
                ][
                    "label"
                ]
            )

            received_by_index[
                criterion_index
            ] = {
                "indice": (
                    criterion_index
                ),

                "criterio": (
                    original_requirement
                ),

                "estado": (
                    state
                ),

                "evidencia": str(
                    item.get(
                        "evidencia"
                    )
                    or ""
                ).strip(),
            }

        # ====================================================
        # 2. REVISAR FALTANTES
        # ====================================================

        missing_indexes = [
            index

            for index
            in range(
                1,
                total_requirements + 1,
            )

            if (
                index
                not in received_by_index
            )
        ]

        if missing_indexes:

            missing_labels = [
                requirements[
                    index - 1
                ][
                    "label"
                ]

                for index
                in missing_indexes
            ]

            raise WeekContentAnalysisError(
                "La respuesta de IA está incompleta. "
                "Faltan requisitos: "
                + ", ".join(
                    missing_labels
                )
            )

        # ====================================================
        # 3. ORDEN ORIGINAL
        # ====================================================

        ordered_criteria = [
            received_by_index[
                index
            ]

            for index
            in range(
                1,
                total_requirements + 1,
            )
        ]

        # ====================================================
        # 4. OBSERVACIÓN PRESENTACIÓN
        # ====================================================

        observation = str(
            payload.get(
                "observacion_presentacion"
            )
            or ""
        ).strip()

        # Si la IA sí devolvió los criterios, pero olvidó construir
        # el texto final, NO desperdiciamos todo el análisis.
        #
        # Lo reconstruimos nosotros a partir de las decisiones
        # realizadas por la IA.
        if not observation:

            observation = (
                self
                ._build_observation_from_criteria(
                    ordered_criteria
                )
            )

        # ====================================================
        # 5. OPCIONALES
        # ====================================================

        optional_files = (
            optional_files
            or {}
        )

        optional_output: Dict[
            str,
            str,
        ] = {}

        optional_labels = {
            "lectura": (
                "Lectura"
            ),

            "video": (
                "Video"
            ),

            "actividad_practica": (
                "Actividad práctica"
            ),

            "ejemplo": (
                "Ejemplo demostrativo"
            ),
        }

        for key in (
            "lectura",
            "video",
            "actividad_practica",
            "ejemplo",
        ):

            files = (
                optional_files.get(
                    key,
                    [],
                )
            )

            # ----------------------------------------------
            # NO EXISTE
            # ----------------------------------------------

            if not files:

                optional_output[
                    key
                ] = ""

                continue

            # ----------------------------------------------
            # EXISTE
            #
            # No necesitamos depender de que la IA escriba
            # correctamente algo que el BACKEND ya sabe.
            # ----------------------------------------------

            names = [
                str(
                    file.get(
                        "name"
                    )
                    or ""
                ).strip()

                for file
                in files

                if (
                    file
                    and str(
                        file.get(
                            "name"
                        )
                        or ""
                    ).strip()
                )
            ]

            optional_output[
                key
            ] = (
                f"{optional_labels[key]} encontrado: "
                + ", ".join(
                    names
                )
            )

        return {
            "criterios": (
                ordered_criteria
            ),

            "observacion_presentacion": (
                observation
            ),

            **optional_output,
        }

    # ========================================================
    # EJECUTAR IA
    # ========================================================

    def analyze_presentation_with_ai(
        self,
        course: CourseCatalog,
        section_label: str,
        requirements: List[
            Dict[str, Any]
        ],
        presentation_info: Dict[
            str,
            Any,
        ],
        presentation_sheet_info: Optional[
            Dict[str, Any]
        ],
        extraction: Dict[
            str,
            Any,
        ],
        optional_files: Optional[
            Dict[
                str,
                List[
                    Dict[str, Any]
                ],
            ]
        ] = None,
        strengthening_info: Optional[
            Dict[str, Any]
        ] = None,
    ) -> Dict[str, Any]:
        """
        Analizar una presentación con IA.

        Si la respuesta estructurada sale incompleta,
        se realiza un segundo intento antes de considerar
        fallida la sección.

        Una falla aquí será atrapada después por analyze_course(),
        por lo que NO detendrá el resto de las semanas.
        """

        call_ai = getattr(
            curriculum_feedback_service,
            "call_ai",
            None,
        )

        if not callable(
            call_ai
        ):

            raise WeekContentAnalysisError(
                "CurriculumFeedbackService "
                "no tiene el método público call_ai()"
            )

        (
            system_message,
            user_message,
        ) = (
            self._build_ai_prompt(
                course=course,

                section_label=(
                    section_label
                ),

                requirements=(
                    requirements
                ),

                presentation_info=(
                    presentation_info
                ),

                presentation_sheet_info=(
                    presentation_sheet_info
                ),

                extraction=(
                    extraction
                ),

                optional_files=(
                    optional_files
                ),

                strengthening_info=(
                    strengthening_info
                ),
            )
        )

        response_schema = (
            self._ai_response_schema(
                requirements
            )
        )

        last_error: Optional[
            Exception
        ] = None

        # ====================================================
        # DOS INTENTOS
        # ====================================================

        for attempt in range(
            1,
            3,
        ):

            print(
                "🤖 [WEEK ANALYSIS] "
                f"IA | sección={section_label} | "
                f"intento={attempt}/2",
                flush=True,
            )

            try:

                current_user_message = (
                    user_message
                )

                # --------------------------------------------
                # SEGUNDO INTENTO
                # --------------------------------------------

                if (
                    attempt
                    == 2
                ):

                    current_user_message += """

============================================================
CORRECCIÓN OBLIGATORIA DE FORMATO
============================================================

La respuesta anterior no pudo validarse completamente.

Debes devolver TODOS los criterios indicados.

Cada criterio debe incluir obligatoriamente:

- indice
- criterio
- estado
- evidencia

El índice debe corresponder EXACTAMENTE con la numeración
de los requisitos proporcionados.

No omitas ninguno.

No cambies la cantidad de criterios.

No respondas con texto fuera del JSON.
"""

                ai_result = (
                    call_ai(
                        system_message=(
                            system_message
                        ),

                        user_message=(
                            current_user_message
                        ),

                        response_schema=(
                            response_schema
                        ),
                    )
                )

                data = (
                    ai_result.get(
                        "data"
                    )
                )

                if not isinstance(
                    data,
                    dict,
                ):

                    raise WeekContentAnalysisError(
                        "El proveedor de IA no devolvió "
                        "un objeto JSON válido"
                    )

                validated = (
                    self._validate_ai_payload(
                        data,
                        requirements,
                        optional_files=(
                            optional_files
                        ),
                    )
                )

                # --------------------------------------------
                # RECONSTRUIR OBSERVACIÓN
                #
                # Lo hacemos nosotros para tener formato
                # consistente entre proveedores.
                # --------------------------------------------

                sheet_found = (
                    None
                )

                if (
                    presentation_sheet_info
                    is not None
                ):

                    sheet_found = bool(
                        presentation_sheet_info.get(
                            "valid_format"
                        )
                    )

                validated[
                    "observacion_presentacion"
                ] = (
                    self
                    ._build_observation_from_criteria(
                        validated[
                            "criterios"
                        ],
                        sheet_found=(
                            sheet_found
                        ),
                    )
                )

                return {
                    "analysis": (
                        validated
                    ),

                    "provider": (
                        ai_result.get(
                            "provider"
                        )
                    ),

                    "model": (
                        ai_result.get(
                            "model"
                        )
                    ),

                    "usage": (
                        ai_result.get(
                            "usage",
                            {},
                        )
                    ),

                    "attempt": (
                        attempt
                    ),
                }

            except Exception as exc:

                last_error = (
                    exc
                )

                print(
                    "⚠️ [WEEK ANALYSIS] "
                    "Respuesta de IA no válida | "
                    f"sección={section_label} | "
                    f"intento={attempt}/2 | "
                    f"tipo={type(exc).__name__} | "
                    f"error={exc}",
                    flush=True,
                )

        raise WeekContentAnalysisError(
            "No fue posible obtener una respuesta "
            "completa de IA para "
            f"{section_label} después de 2 intentos. "
            f"Último error: {last_error}"
        )

    # ========================================================
    # ESCRIBIR GOOGLE SHEETS
    # ========================================================

    def _write_google_sheet_matrix(
        self,
        matrix_file_id: str,
        sheet_title: str,
        row_values: Dict[
            int,
            str,
        ],
    ) -> Dict[str, Any]:

        escaped_title = (
            sheet_title.replace(
                "'",
                "''",
            )
        )

        updates = [
            {
                "range": (
                    f"'{escaped_title}'!"
                    f"{MATRIX_AI_COLUMN}"
                    f"{row}"
                ),

                "values": [
                    [
                        value
                    ]
                ],
            }

            for row, value
            in sorted(
                row_values.items()
            )
        ]

        result = (
            google_sheets_service
            .batch_update_values(
                matrix_file_id,
                updates,
            )
        )

        return {
            "success": True,

            "mode": (
                "google_sheets_api"
            ),

            "sheet": (
                sheet_title
            ),

            "column": (
                MATRIX_AI_COLUMN
            ),

            "updated": (
                result.get(
                    "updated",
                    0,
                )
            ),

            "targets": (
                len(
                    row_values
                )
            ),
        }


    # ========================================================
    # ESCRIBIR XLSX
    # ========================================================

    def _write_xlsx_matrix(
        self,
        matrix_file_id: str,
        sheet_title: str,
        row_values: Dict[
            int,
            str,
        ],
    ) -> Dict[str, Any]:

        content = (
            drive_service
            .download_file(
                matrix_file_id
            )
        )

        if not content:

            raise WeekContentAnalysisError(
                "No se pudo descargar "
                "la matriz XLSX para "
                "escribir resultados"
            )

        workbook = None

        try:

            workbook = (
                openpyxl
                .load_workbook(
                    io.BytesIO(
                        content
                    ),
                    read_only=False,
                    data_only=False,
                    keep_links=True,
                )
            )

            if (
                sheet_title
                not in workbook.sheetnames
            ):

                raise WeekContentAnalysisError(
                    "No existe la hoja "
                    f"'{sheet_title}' "
                    "en la matriz XLSX"
                )

            worksheet = (
                workbook[
                    sheet_title
                ]
            )

            # IMPORTANTE:
            #
            # En la matriz las observaciones de Presentación
            # y Lectura están combinadas verticalmente.
            #
            # Nosotros ya detectamos dinámicamente la PRIMERA
            # fila del bloque, que es la celda superior izquierda
            # de la combinación.
            #
            # Por eso sí se puede asignar directamente H{fila}.

            for row, value in (
                row_values.items()
            ):

                worksheet[
                    f"{MATRIX_AI_COLUMN}"
                    f"{row}"
                ] = (
                    value
                )

            output = (
                io.BytesIO()
            )

            workbook.save(
                output
            )

            updated = (
                drive_service
                .upload_file(
                    output.getvalue(),
                    XLSX_MIME,
                    matrix_file_id,
                )
            )

            if not updated:

                raise WeekContentAnalysisError(
                    "Google Drive no pudo "
                    "actualizar la matriz XLSX"
                )

            return {
                "success": True,

                "mode": (
                    "xlsx_upload"
                ),

                "sheet": (
                    sheet_title
                ),

                "column": (
                    MATRIX_AI_COLUMN
                ),

                "updated": (
                    len(
                        row_values
                    )
                ),

                "targets": (
                    len(
                        row_values
                    )
                ),
            }

        finally:

            if workbook is not None:
                workbook.close()


    def write_matrix(
        self,
        matrix: Dict[
            str,
            Any,
        ],
        row_values: Dict[
            int,
            str,
        ],
    ) -> Dict[str, Any]:

        matrix_file = (
            matrix[
                "file"
            ]
        )

        sheet_title = (
            matrix[
                "sheet_title"
            ]
        )

        # ----------------------------------------------------
        # GOOGLE SHEETS
        # ----------------------------------------------------

        if (
            matrix_file.get(
                "mimeType"
            )
            == GOOGLE_SHEET_MIME
        ):

            return (
                self._write_google_sheet_matrix(
                    matrix_file[
                        "id"
                    ],
                    sheet_title,
                    row_values,
                )
            )

        # ----------------------------------------------------
        # XLSX
        # ----------------------------------------------------

        if (
            matrix_file.get(
                "mimeType"
            )
            == XLSX_MIME
        ):

            return (
                self._write_xlsx_matrix(
                    matrix_file[
                        "id"
                    ],
                    sheet_title,
                    row_values,
                )
            )

        raise WeekContentAnalysisError(
            "Formato de matriz "
            "no soportado para escritura"
        )


    # ========================================================
    # ANALIZAR CURSO COMPLETO
    # ========================================================

    def analyze_course(
        self,
        course: CourseCatalog,
        course_folder_id: str,
        course_folder_name: str,
        area_folder_name: str,
        write_output: bool = True,
        progress_callback: Optional[
            Callable[
                [
                    int,
                    str,
                    Optional[int],
                ],
                None,
            ]
        ] = None,
    ) -> Dict[str, Any]:
        """
        Analizar UN curso.

        Se realizan 11 análisis independientes:

        1. Semana Diagnóstico
        2. Semana 2
        3. Semana 3
        ...
        11. Semana 11

        Esto permite:
        - menor contexto por llamada;
        - mejor precisión;
        - progreso real del job;
        - aislar cada semana;
        - evitar que contenido de una semana contamine otra.
        """

        started_at = (
            time.perf_counter()
        )


        def progress(
            value: int,
            stage: str,
            week: Optional[
                int
            ] = None,
        ) -> None:

            print(
                "📊 [WEEK ANALYSIS] "
                f"{value}% | "
                f"{stage}",
                flush=True,
            )

            if progress_callback:

                progress_callback(
                    value,
                    stage,
                    week,
                )


        # ====================================================
        # 1. VALIDAR CURSO Y PERÍODO
        # ====================================================

        progress(
            10,
            "Validando curso y período",
            None,
        )

        period = (
            self.validate_course_identity(
                course,
                course_folder_name,
                area_folder_name,
            )
        )

        code = str(
            course.code
        )

        target_period = (
            period[
                "target_period"
            ]
        )


        # ====================================================
        # 2. MATRIZ
        # ====================================================

        progress(
            14,
            "Leyendo F2_Semanas",
            None,
        )

        matrix = (
            self.load_matrix(
                course_folder_id
            )
        )

        blocks = (
            matrix[
                "blocks"
            ]
        )


        # ====================================================
        # 3. LOCALIZAR MATERIALES
        # ====================================================

        progress(
            17,
            "Localizando Diagnóstico y Contenidos",
            None,
        )

        diagnostic = (
            self.inspect_diagnostic(
                course_folder_id,
                code,
                target_period,
            )
        )

        contents_folder = (
            self._find_single_folder(
                course_folder_id,
                CONTENTS_FOLDER_NAME,
                required=False,
            )
        )

        weeks = {
            week: (
                self.inspect_week(
                    contents_folder,
                    code,
                    target_period,
                    week,
                )
            )

            for week
            in range(
                2,
                12,
            )
        }


        # ====================================================
        # RESULTADOS EN MEMORIA
        # ====================================================

        row_values: Dict[
            int,
            str,
        ] = {}

        providers_used: List[
            Dict[str, Any]
        ] = []

        results: Dict[
            str,
            Any,
        ] = {}

        warnings: List[
            str
        ] = []


        # ====================================================
        # 4. SEMANA DIAGNÓSTICO
        # ====================================================

        progress(
            20,
            "Analizando Semana Diagnóstico",
            None,
        )

        diagnostic_block = (
            blocks[
                "diagnostic"
            ]
        )

        try:

            diagnostic_extraction = (
                self.extract_presentation_text(
                    diagnostic[
                        "presentation"
                    ][
                        "file"
                    ]
                )
            )

            if diagnostic_extraction.get(
                "truncated"
            ):

                warnings.append(
                    "Semana Diagnóstico: "
                    "el texto extraído del PDF "
                    "fue recortado"
                )

            if not diagnostic_extraction.get(
                "success"
            ):

                warnings.append(
                    "Semana Diagnóstico: "
                    + str(
                        diagnostic_extraction.get(
                            "error"
                        )
                        or
                        "no se pudo leer el PDF"
                    )
                )

            diagnostic_ai = (
                self.analyze_presentation_with_ai(
                    course=course,

                    section_label=(
                        diagnostic_block[
                            "label"
                        ]
                    ),

                    requirements=(
                        diagnostic_block[
                            "presentation"
                        ][
                            "requirements"
                        ]
                    ),

                    presentation_info=(
                        diagnostic[
                            "presentation"
                        ]
                    ),

                    presentation_sheet_info=(
                        None
                    ),

                    extraction=(
                        diagnostic_extraction
                    ),

                    optional_files=(
                        None
                    ),

                    strengthening_info=(
                        diagnostic[
                            "strengthening"
                        ]
                    ),
                )
            )

            providers_used.append(
                {
                    "section": (
                        "Semana Diagnóstico"
                    ),

                    "provider": (
                        diagnostic_ai.get(
                            "provider"
                        )
                    ),

                    "model": (
                        diagnostic_ai.get(
                            "model"
                        )
                    ),

                    "usage": (
                        diagnostic_ai.get(
                            "usage",
                            {},
                        )
                    ),

                    "attempt": (
                        diagnostic_ai.get(
                            "attempt",
                            1,
                        )
                    ),
                }
            )

            row_values[
                diagnostic_block[
                    "presentation"
                ][
                    "row"
                ]
            ] = (
                diagnostic_ai[
                    "analysis"
                ][
                    "observacion_presentacion"
                ]
            )

            results[
                "diagnostic"
            ] = {
                "success": True,

                "files": (
                    diagnostic
                ),

                "extraction": {
                    key: value

                    for key, value
                    in diagnostic_extraction.items()

                    if key != "text"
                },

                "analysis": (
                    diagnostic_ai[
                        "analysis"
                    ]
                ),
            }

        except Exception as exc:

            error_message = str(
                exc
            )

            print(
                "❌ [WEEK ANALYSIS] "
                "Semana Diagnóstico falló, "
                "pero el curso continuará | "
                f"tipo={type(exc).__name__} | "
                f"error={error_message}",
                flush=True,
            )

            warnings.append(
                "Semana Diagnóstico: "
                "no se pudo completar el análisis "
                f"con IA. {error_message}"
            )

            results[
                "diagnostic"
            ] = {
                "success": False,

                "files": (
                    diagnostic
                ),

                "error": (
                    error_message
                ),
            }

        # ====================================================
        # 5. SEMANAS 2..11
        # ====================================================

        week_progress = {
            2: 27,
            3: 34,
            4: 41,
            5: 48,
            6: 55,
            7: 62,
            8: 69,
            9: 76,
            10: 83,
            11: 90,
        }

        successful_weeks = 0
        failed_weeks = 0

        for week in range(
            2,
            12,
        ):

            progress(
                week_progress[
                    week
                ],
                f"Analizando Semana {week}",
                week,
            )

            block = (
                blocks[
                    f"week_{week}"
                ]
            )

            week_info = (
                weeks[
                    week
                ]
            )

            try:

                # =============================================
                # 1. SOLO PDF
                # =============================================

                extraction = (
                    self.extract_presentation_text(
                        week_info[
                            "presentation"
                        ][
                            "file"
                        ]
                    )
                )

                if extraction.get(
                    "truncated"
                ):

                    warnings.append(
                        f"Semana {week}: "
                        "el texto extraído del PDF "
                        "fue recortado por el límite "
                        "interno del backend"
                    )

                if not extraction.get(
                    "success"
                ):

                    warnings.append(
                        f"Semana {week}: "
                        + str(
                            extraction.get(
                                "error"
                            )
                            or
                            "no se pudo leer el PDF"
                        )
                    )

                # =============================================
                # 2. IA
                # =============================================

                ai = (
                    self.analyze_presentation_with_ai(
                        course=course,

                        section_label=(
                            block[
                                "label"
                            ]
                        ),

                        requirements=(
                            block[
                                "presentation"
                            ][
                                "requirements"
                            ]
                        ),

                        presentation_info=(
                            week_info[
                                "presentation"
                            ]
                        ),

                        presentation_sheet_info=(
                            week_info[
                                "presentation_sheet"
                            ]
                        ),

                        extraction=(
                            extraction
                        ),

                        optional_files=(
                            week_info[
                                "optional"
                            ]
                        ),

                        strengthening_info=(
                            None
                        ),
                    )
                )

                analysis = (
                    ai[
                        "analysis"
                    ]
                )

                # =============================================
                # 3. PROVIDER
                # =============================================

                providers_used.append(
                    {
                        "section": (
                            f"Semana {week}"
                        ),

                        "provider": (
                            ai.get(
                                "provider"
                            )
                        ),

                        "model": (
                            ai.get(
                                "model"
                            )
                        ),

                        "usage": (
                            ai.get(
                                "usage",
                                {},
                            )
                        ),

                        "attempt": (
                            ai.get(
                                "attempt",
                                1,
                            )
                        ),
                    }
                )

                # =============================================
                # 4. PRESENTACIÓN
                # =============================================

                presentation_row = (
                    block[
                        "presentation"
                    ][
                        "row"
                    ]
                )

                row_values[
                    presentation_row
                ] = (
                    analysis[
                        "observacion_presentacion"
                    ]
                )

                # =============================================
                # 5. LECTURA
                # =============================================

                lectura_row = (
                    block[
                        "lectura"
                    ][
                        "row"
                    ]
                )

                if (
                    lectura_row
                    is not None
                ):

                    row_values[
                        lectura_row
                    ] = (
                        analysis.get(
                            "lectura",
                            "",
                        )
                    )

                # =============================================
                # 6. VARIOS
                # =============================================

                for key in (
                    "video",
                    "actividad_practica",
                    "ejemplo",
                ):

                    target_row = (
                        block[
                            "varios"
                        ].get(
                            key
                        )
                    )

                    if (
                        target_row
                        is not None
                    ):

                        row_values[
                            target_row
                        ] = (
                            analysis.get(
                                key,
                                "",
                            )
                        )

                # =============================================
                # 7. RESULTADO
                # =============================================

                results[
                    f"week_{week}"
                ] = {
                    "success": True,

                    "files": (
                        week_info
                    ),

                    "extraction": {
                        key: value

                        for key, value
                        in extraction.items()

                        if (
                            key != "text"
                        )
                    },

                    "analysis": (
                        analysis
                    ),
                }

                successful_weeks += 1

                print(
                    "✅ [WEEK ANALYSIS] "
                    f"Semana {week} completada",
                    flush=True,
                )

            # =================================================
            # ERROR DE ESTA SEMANA
            #
            # NO PROPAGAR.
            # NO DETENER EL CURSO.
            # =================================================

            except Exception as exc:

                failed_weeks += 1

                error_message = str(
                    exc
                )

                print(
                    "❌ [WEEK ANALYSIS] "
                    f"Semana {week} falló, "
                    "pero el curso continuará | "
                    f"tipo={type(exc).__name__} | "
                    f"error={error_message}",
                    flush=True,
                )

                warnings.append(
                    f"Semana {week}: "
                    "no se pudo completar el análisis "
                    f"con IA. {error_message}"
                )

                # ---------------------------------------------
                # IMPORTANTE:
                #
                # No escribimos una observación curricular
                # inventada.
                #
                # Sí dejamos constancia en el RESULT JSON.
                #
                # Y NO tocamos la celda H de Presentación para
                # evitar borrar una observación anterior válida.
                # ---------------------------------------------

                results[
                    f"week_{week}"
                ] = {
                    "success": False,

                    "files": (
                        week_info
                    ),

                    "error": (
                        error_message
                    ),
                }

                # ---------------------------------------------
                # OPCIONALES
                #
                # Estos NO dependen de IA.
                # Podemos seguir actualizándolos aunque haya
                # fallado el análisis del PDF.
                # ---------------------------------------------

                optional_files = (
                    week_info.get(
                        "optional",
                        {}
                    )
                )

                lectura_row = (
                    block[
                        "lectura"
                    ][
                        "row"
                    ]
                )

                if (
                    lectura_row
                    is not None
                ):

                    lectura_files = (
                        optional_files.get(
                            "lectura",
                            [],
                        )
                    )

                    if lectura_files:

                        names = [
                            str(
                                item.get(
                                    "name"
                                )
                                or ""
                            )

                            for item
                            in lectura_files
                        ]

                        row_values[
                            lectura_row
                        ] = (
                            "Lectura encontrada: "
                            + ", ".join(
                                names
                            )
                        )

                    else:

                        row_values[
                            lectura_row
                        ] = ""

                optional_labels = {
                    "video": (
                        "Video"
                    ),

                    "actividad_practica": (
                        "Actividad práctica"
                    ),

                    "ejemplo": (
                        "Ejemplo demostrativo"
                    ),
                }

                for key in (
                    "video",
                    "actividad_practica",
                    "ejemplo",
                ):

                    target_row = (
                        block[
                            "varios"
                        ].get(
                            key
                        )
                    )

                    if (
                        target_row
                        is None
                    ):

                        continue

                    files = (
                        optional_files.get(
                            key,
                            [],
                        )
                    )

                    if files:

                        names = [
                            str(
                                item.get(
                                    "name"
                                )
                                or ""
                            )

                            for item
                            in files
                        ]

                        row_values[
                            target_row
                        ] = (
                            f"{optional_labels[key]} "
                            "encontrado: "
                            + ", ".join(
                                names
                            )
                        )

                    else:

                        row_values[
                            target_row
                        ] = ""

                # MUY IMPORTANTE:
                #
                # continúa Semana 7, 8, 9...
                continue


        # ====================================================
        # 6. ESCRIBIR MATRIZ
        # ====================================================

        write_result = (
            None
        )

        if write_output:

            progress(
                96,
                (
                    "Escribiendo Observaciones IA "
                    "en F2_Semanas"
                ),
                None,
            )

            write_result = (
                self.write_matrix(
                    matrix,
                    row_values,
                )
            )


        # ====================================================
        # 7. FINAL
        # ====================================================

        progress(
            99,
            "Finalizando análisis",
            None,
        )
        
        sections_successful = sum(
            1

            for item
            in results.values()

            if (
                item.get(
                    "success"
                )
                is True
            )
        )

        sections_failed = sum(
            1

            for item
            in results.values()

            if (
                item.get(
                    "success"
                )
                is False
            )
        )

        return {
            "success": True,

            "course": {
                "code": (
                    code
                ),

                "name": (
                    course.name
                ),

                "area": (
                    course_contacts_service
                    .canonical_area(
                        course.area
                    )
                ),

                "folder_id": (
                    course_folder_id
                ),

                "folder_name": (
                    course_folder_name
                ),

                "area_folder_name": (
                    area_folder_name
                ),
            },

            "period": (
                period
            ),

            "matrix": {
                "file": (
                    self._file_public_data(
                        matrix[
                            "file"
                        ]
                    )
                ),

                "sheet": (
                    matrix[
                        "sheet_title"
                    ]
                ),

                "column": (
                    MATRIX_AI_COLUMN
                ),

                "updates_prepared": (
                    len(
                        row_values
                    )
                ),
            },

            "warnings": (
                warnings
            ),

            "providers_used": (
                providers_used
            ),

            "results": (
                results
            ),

            "write": (
                write_result
            ),

            "elapsed_seconds": round(
                (
                    time.perf_counter()
                    - started_at
                ),
                2,
            ),

            "summary": {
                "total_sections": (
                    len(
                        results
                    )
                ),

                "successful_sections": (
                    sections_successful
                ),

                "failed_sections": (
                    sections_failed
                ),

                "partial_success": (
                    sections_failed > 0
                ),
            },
        }


week_content_analysis_service = (
    WeekContentAnalysisService()
)